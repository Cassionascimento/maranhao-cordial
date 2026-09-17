"""Presentation Engine (P5.X — Milestone 4).

Gera um arquivo .pptx REAL (biblioteca `python-pptx`, sem serviço
externo, sem custo, sem credencial nova) a partir da ata executiva já
montada pelo Secretário (M3) -- nunca HTML fingindo apresentação, nunca
PDF, nunca texto solto. Persistência via mi_artefatos.registrar_artefato
(M1): este módulo só GERA bytes, nunca decide o que é aprovado nem
publica nada.

Design system: paleta e tipografia reaproveitadas do próprio Admin V2
(maranhao-backend/app-shell.css: fundo #07100b, dourado #b99a5d, creme
#efe7d6) -- o deck usa a identidade visual já em produção da Maranhão
Cordial, não um tema genérico de biblioteca.

3 a 5 slides por padrão (capa, narrativa, decisões, próximos passos) --
mais slides só quando o Chart Engine (M5) anexar gráficos reais."""
import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from mi_artefatos import registrar_artefato
from mi_conselho import AGENTES

MIME_PPTX = 'application/vnd.openxmlformats-officedocument.presentationml.presentation'

# Paleta reaproveitada de maranhao-backend/app-shell.css -- mesma
# identidade do Admin V2, nunca um tema genérico novo.
COR_FUNDO = RGBColor(0x07, 0x10, 0x0B)
COR_DOURADO = RGBColor(0xB9, 0x9A, 0x5D)
COR_CREME = RGBColor(0xEF, 0xE7, 0xD6)
COR_MUTED = RGBColor(0x98, 0xA4, 0x99)

FONTE_TITULO = 'Georgia'
FONTE_CORPO = 'Arial'

LARGURA_SLIDE = Inches(13.333)  # widescreen 16:9
ALTURA_SLIDE = Inches(7.5)

MAX_ITENS_POR_SLIDE_LISTA = 5


def _fundo_solido(slide, cor):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = cor


def _caixa_texto(slide, left, top, width, height, texto, *, tamanho=18, cor=COR_CREME,
                  fonte=FONTE_CORPO, negrito=False, alinhamento=PP_ALIGN.LEFT,
                  ancora=MSO_ANCHOR.TOP):
    caixa = slide.shapes.add_textbox(left, top, width, height)
    quadro = caixa.text_frame
    quadro.word_wrap = True
    quadro.vertical_anchor = ancora
    paragrafo = quadro.paragraphs[0]
    paragrafo.alignment = alinhamento
    corrida = paragrafo.add_run()
    corrida.text = texto
    corrida.font.size = Pt(tamanho)
    corrida.font.color.rgb = cor
    corrida.font.name = fonte
    corrida.font.bold = negrito
    return caixa


def _slide_vazio(apresentacao):
    layout_em_branco = apresentacao.slide_layouts[6]  # "Blank" do template default do python-pptx
    slide = apresentacao.slides.add_slide(layout_em_branco)
    _fundo_solido(slide, COR_FUNDO)
    return slide


def _nome_agente(codigo):
    return AGENTES.get(codigo, {}).get('nome', codigo)


def _slide_capa(apresentacao, ata):
    slide = _slide_vazio(apresentacao)
    _caixa_texto(slide, Inches(0.9), Inches(0.6), Inches(6), Inches(0.4),
                 'MARANHÃO INTELLIGENCE · CONSELHO', tamanho=12, cor=COR_DOURADO, negrito=True)
    _caixa_texto(slide, Inches(0.9), Inches(2.5), Inches(11.5), Inches(2.3),
                 ata.get('titulo') or 'Reunião do Conselho', tamanho=40, cor=COR_CREME, fonte=FONTE_TITULO)
    participantes = ', '.join(_nome_agente(a) for a in (ata.get('participantes') or []))
    _caixa_texto(slide, Inches(0.9), Inches(5.3), Inches(11.5), Inches(0.6),
                 participantes or 'Participantes não identificados nesta reunião.', tamanho=14, cor=COR_MUTED)
    return slide


def _slide_narrativa(apresentacao, ata):
    slide = _slide_vazio(apresentacao)
    _caixa_texto(slide, Inches(0.9), Inches(0.6), Inches(6), Inches(0.4),
                 'CONTEXTO', tamanho=12, cor=COR_DOURADO, negrito=True)
    _caixa_texto(slide, Inches(0.9), Inches(1.8), Inches(11.5), Inches(4.5),
                 ata.get('narrativa') or 'Sem narrativa disponível para esta reunião.',
                 tamanho=26, cor=COR_CREME, fonte=FONTE_TITULO, ancora=MSO_ANCHOR.MIDDLE)
    return slide


def _slide_lista(apresentacao, rotulo, itens, *, cor_numero=COR_DOURADO, texto_vazio):
    """Um dominante-idea-por-slide: números grandes + uma linha por item,
    nunca uma parede de bullets. Nunca omitido silenciosamente quando a
    lista está vazia -- mostra o texto_vazio explícito em vez disso."""
    slide = _slide_vazio(apresentacao)
    _caixa_texto(slide, Inches(0.9), Inches(0.6), Inches(6), Inches(0.4),
                 rotulo, tamanho=12, cor=COR_DOURADO, negrito=True)
    itens = list(itens or [])[:MAX_ITENS_POR_SLIDE_LISTA]
    if not itens:
        _caixa_texto(slide, Inches(0.9), Inches(3.2), Inches(11), Inches(1), texto_vazio,
                     tamanho=20, cor=COR_MUTED)
        return slide
    y = Inches(1.7)
    altura_item = Inches(1.0)
    for indice, item in enumerate(itens, start=1):
        _caixa_texto(slide, Inches(0.9), y, Inches(0.9), altura_item, str(indice),
                     tamanho=32, cor=cor_numero, fonte=FONTE_TITULO, negrito=True)
        _caixa_texto(slide, Inches(2.0), y, Inches(10.3), altura_item, item,
                     tamanho=18, cor=COR_CREME, ancora=MSO_ANCHOR.MIDDLE)
        y = Emu(int(y) + int(altura_item))
    return slide


def montar_deck_executivo(ata):
    """Monta o .pptx a partir da ata já pronta (M3) -- capa, narrativa,
    decisões e próximos passos (3 a 5 slides por padrão). Nunca inventa
    texto: cada slide só reformata o que a ata já trouxe. M5 estende esta
    função para anexar slides de gráfico real quando houver dados
    suficientes."""
    apresentacao = Presentation()
    apresentacao.slide_width = LARGURA_SLIDE
    apresentacao.slide_height = ALTURA_SLIDE

    _slide_capa(apresentacao, ata)
    _slide_narrativa(apresentacao, ata)
    _slide_lista(apresentacao, 'DECISÕES', ata.get('decisoes'), texto_vazio='Nenhuma decisão registrada nesta reunião.')
    _slide_lista(apresentacao, 'PRÓXIMOS PASSOS', ata.get('proximos_passos'), cor_numero=COR_CREME,
                 texto_vazio='Nenhum próximo passo registrado nesta reunião.')

    buffer = io.BytesIO()
    apresentacao.save(buffer)
    return buffer.getvalue()


def gerar_e_registrar_apresentacao(factory, ata, *, agent_id=None, decision_id=None, parent_artifact_id=None):
    """Gera o .pptx e persiste como artefato (M1) -- nasce em 'gerado',
    só vira 'aprovado' por ação humana explícita (mi_artefatos.
    aprovar_artefato), nunca publicado sozinho."""
    conteudo = montar_deck_executivo(ata)
    meeting_id = ata.get('meeting_id')
    return registrar_artefato(
        factory, artifact_type='PRESENTATION', conteudo=conteudo, mime_type=MIME_PPTX,
        source_type='conselho', source_id=meeting_id, meeting_id=meeting_id,
        agent_id=agent_id, decision_id=decision_id, parent_artifact_id=parent_artifact_id,
        metadata={'titulo': ata.get('titulo'), 'participantes': ata.get('participantes')},
    )


def registrar_rotas(app, factory, autorizado):
    from flask import jsonify
    from psycopg2.extras import RealDictCursor

    from mi_conselho import buscar_registro
    from mi_secretario_executivo import montar_ata_executiva

    @app.route('/api/admin/mi/conselho/<registro_id>/apresentacao', methods=['POST'])
    def mi_presentation_gerar(registro_id):
        if not autorizado():
            return jsonify(success=False, error='Não autorizado.'), 401
        conn = factory()
        try:
            conn.set_session(readonly=True, isolation_level='REPEATABLE READ')
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                cur.execute("SET LOCAL statement_timeout='10s'")
                registro = buscar_registro(cur, registro_id)
        except Exception:
            app.logger.exception('Falha ao buscar registro para montar a apresentação')
            return jsonify(success=False, error='Apresentação indisponível.'), 503
        finally:
            conn.close()
        if not registro:
            return jsonify(success=False, error='Reunião/registro não encontrado.'), 404
        try:
            ata = montar_ata_executiva(registro)
            resultado = gerar_e_registrar_apresentacao(factory, ata)
        except Exception:
            app.logger.exception('Falha ao gerar apresentação executiva')
            return jsonify(success=False, error='Não foi possível gerar a apresentação.'), 503
        sucesso = resultado.pop('success', False)
        return jsonify(success=sucesso, **resultado), (201 if sucesso else 503)
