"""Chart Engine (P5.X — Milestone 5).

Especificação canônica de gráfico, reutilizável entre o Command Center
(imagem SVG standalone) e o Presentation Engine (gráfico nativo do
PowerPoint, M4). Gráficos são SEMPRE gerados programaticamente a partir
de números reais -- nunca pede a um modelo de imagem para "desenhar" um
gráfico factual, e nunca fabrica dado quando não há o suficiente: nesse
caso o chamador recebe NOT_ENOUGH_DATA (mesmo vocabulário de proveniência
do contrato P5 -- mi_intelligence_api.py) em vez de um gráfico.

Nenhuma dependência nova: SVG é texto (stdlib), gráfico nativo de PPTX
usa python-pptx (já adicionado no M4)."""
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from mi_artefatos import registrar_artefato

TIPOS_GRAFICO = ('bar', 'line', 'pie')
PROVENIENCIA_VALIDA = ('REAL', 'DERIVED', 'INFERRED', 'SYNTHETIC_TEST', 'NOT_ENOUGH_DATA')
NOT_ENOUGH_DATA = 'NOT_ENOUGH_DATA'
MIME_SVG = 'image/svg+xml'

_TIPO_PPTX = {
    'bar': XL_CHART_TYPE.COLUMN_CLUSTERED,
    'line': XL_CHART_TYPE.LINE,
    'pie': XL_CHART_TYPE.PIE,
}

# Paleta reaproveitada de maranhao-backend/app-shell.css e de
# mi_presentation_engine.py -- mesma identidade visual em toda a
# superfície criativa, nunca um tema genérico.
_COR_FUNDO = '#07100b'
_COR_DOURADO = '#b99a5d'
_COR_CREME = '#efe7d6'
_COR_MUTED = '#98a49b'


def validar_chart_spec(spec):
    """Valida a forma do contrato canônico -- nunca aceita um gráfico com
    séries de tamanho inconsistente ou proveniência fora do vocabulário
    já estabelecido no P5."""
    if not isinstance(spec, dict):
        raise ValueError('chart_spec_invalido')
    if spec.get('chart_type') not in TIPOS_GRAFICO:
        raise ValueError('chart_type_invalido')
    x = spec.get('x')
    if not isinstance(x, list) or not x:
        raise ValueError('eixo_x_obrigatorio')
    series = spec.get('series')
    if not isinstance(series, list) or not series:
        raise ValueError('series_obrigatoria')
    if spec['chart_type'] in ('bar','pie') and len(series)!=1:
        raise ValueError('tipo_exige_serie_unica')
    if len(x)>50 or len(series)>8:
        raise ValueError('grafico_excede_limite')
    import math
    for serie in series:
        if not isinstance(serie, dict) or 'name' not in serie or 'values' not in serie:
            raise ValueError('serie_invalida')
        if not isinstance(serie['values'],list) or any(isinstance(v,bool) or not isinstance(v,(int,float)) or not math.isfinite(v) or v<0 for v in serie['values']):
            raise ValueError('valores_nao_suportados')
        if len(serie['values']) != len(x):
            raise ValueError('serie_com_tamanho_diferente_do_eixo_x')
    if spec.get('confidence') not in PROVENIENCIA_VALIDA:
        raise ValueError('confidence_invalida')
    if not (spec.get('title') or '').strip():
        raise ValueError('titulo_obrigatorio')
    return spec


def montar_chart_spec_ou_aguardando(*, chart_type, title, x, series, units, source, freshness,
                                     confidence, minimo_pontos=1):
    """Ponto único de decisão 'temos dado suficiente?' -- nunca deixa o
    chamador decidir isso caso a caso. Sem pontos suficientes (learning/
    território/forecast ainda não maduros, por exemplo), devolve
    NOT_ENOUGH_DATA explicitamente em vez de um chart_spec, para o
    consumidor (Admin/PPTX) mostrar 'Aguardando dados' -- nunca um
    gráfico vazio ou inventado."""
    if confidence == NOT_ENOUGH_DATA or not x or len(x) < minimo_pontos or not series:
        return NOT_ENOUGH_DATA
    spec = {
        'chart_type': chart_type, 'title': title, 'x': list(x), 'series': series,
        'units': units, 'source': source, 'freshness': freshness, 'confidence': confidence,
    }
    return validar_chart_spec(spec)


def _escala(valores, altura_util):
    maximo = max((abs(v) for serie_valores in valores for v in serie_valores), default=0)
    if maximo <= 0:
        return lambda v: 0
    return lambda v: (v / maximo) * altura_util


def _svg_bar(spec, largura, altura):
    margem_esq, margem_topo, margem_baixo = 70, 60, 50
    largura_util = largura - margem_esq - 30
    altura_util = altura - margem_topo - margem_baixo
    serie = spec['series'][0]
    escala = _escala([serie['values']], altura_util)
    n = len(spec['x'])
    largura_barra = largura_util / n * 0.6
    passo = largura_util / n
    partes = []
    for i, (rotulo, valor) in enumerate(zip(spec['x'], serie['values'])):
        h = escala(valor)
        x0 = margem_esq + i * passo + (passo - largura_barra) / 2
        y0 = margem_topo + altura_util - h
        partes.append(f'<rect x="{x0:.1f}" y="{y0:.1f}" width="{largura_barra:.1f}" height="{h:.1f}" fill="{_COR_DOURADO}" />')
        partes.append(f'<text x="{x0 + largura_barra / 2:.1f}" y="{y0 - 8:.1f}" fill="{_COR_CREME}" font-size="13" text-anchor="middle">{valor:g}</text>')
        partes.append(f'<text x="{x0 + largura_barra / 2:.1f}" y="{margem_topo + altura_util + 22:.1f}" fill="{_COR_MUTED}" font-size="12" text-anchor="middle">{rotulo}</text>')
    partes.append(f'<line x1="{margem_esq}" y1="{margem_topo + altura_util:.1f}" x2="{largura - 30}" y2="{margem_topo + altura_util:.1f}" stroke="{_COR_MUTED}" stroke-width="1" />')
    return ''.join(partes)


def _svg_line(spec, largura, altura):
    margem_esq, margem_topo, margem_baixo = 70, 60, 50
    largura_util = largura - margem_esq - 30
    altura_util = altura - margem_topo - margem_baixo
    n = len(spec['x'])
    passo = largura_util / max(n - 1, 1)
    partes = []
    todas_series = [s['values'] for s in spec['series']]
    escala = _escala(todas_series, altura_util)
    cores = [_COR_DOURADO, _COR_CREME, _COR_MUTED]
    for idx_serie, serie in enumerate(spec['series']):
        cor = cores[idx_serie % len(cores)]
        pontos = []
        for i, valor in enumerate(serie['values']):
            x = margem_esq + i * passo
            y = margem_topo + altura_util - escala(valor)
            pontos.append(f'{x:.1f},{y:.1f}')
        partes.append(f'<polyline points="{" ".join(pontos)}" fill="none" stroke="{cor}" stroke-width="2.5" />')
        for i, valor in enumerate(serie['values']):
            x = margem_esq + i * passo
            y = margem_topo + altura_util - escala(valor)
            partes.append(f'<circle cx="{x:.1f}" cy="{y:.1f}" r="3.5" fill="{cor}" />')
    for i, rotulo in enumerate(spec['x']):
        x = margem_esq + i * passo
        partes.append(f'<text x="{x:.1f}" y="{margem_topo + altura_util + 22:.1f}" fill="{_COR_MUTED}" font-size="12" text-anchor="middle">{rotulo}</text>')
    partes.append(f'<line x1="{margem_esq}" y1="{margem_topo + altura_util:.1f}" x2="{largura - 30}" y2="{margem_topo + altura_util:.1f}" stroke="{_COR_MUTED}" stroke-width="1" />')
    return ''.join(partes)


def _svg_pie(spec, largura, altura):
    import math
    serie = spec['series'][0]
    total = sum(serie['values']) or 1
    cx, cy, raio = largura / 2, altura / 2 + 10, min(largura, altura) / 3
    cores = [_COR_DOURADO, _COR_CREME, _COR_MUTED, '#6f8175', '#b7c4ba']
    partes = []
    angulo = -math.pi / 2
    for i, (rotulo, valor) in enumerate(zip(spec['x'], serie['values'])):
        fatia = (valor / total) * 2 * math.pi
        x1, y1 = cx + raio * math.cos(angulo), cy + raio * math.sin(angulo)
        angulo += fatia
        x2, y2 = cx + raio * math.cos(angulo), cy + raio * math.sin(angulo)
        grande_arco = 1 if fatia > math.pi else 0
        cor = cores[i % len(cores)]
        partes.append(
            f'<path d="M{cx:.1f},{cy:.1f} L{x1:.1f},{y1:.1f} A{raio:.1f},{raio:.1f} 0 {grande_arco} 1 {x2:.1f},{y2:.1f} Z" fill="{cor}" />'
        )
        partes.append(f'<text x="{10}" y="{altura - 15 - i * 18}" fill="{_COR_CREME}" font-size="12">{rotulo}: {valor:g}</text>')
    return ''.join(partes)


_RENDERIZADORES_SVG = {'bar': _svg_bar, 'line': _svg_line, 'pie': _svg_pie}


def renderizar_svg(spec, largura=640, altura=400):
    """Gráfico factual como SVG standalone (artefato IMAGE/CHART) -- texto
    puro (stdlib), sem nenhuma dependência de rasterização."""
    validar_chart_spec(spec)
    from html import escape
    spec = {**spec, 'x':[escape(str(x),quote=True) for x in spec['x']],
            'title':escape(str(spec['title']),quote=True),
            'source':escape(str(spec.get('source') or '—'),quote=True),
            'freshness':escape(str(spec.get('freshness') or ''),quote=True)}
    corpo = _RENDERIZADORES_SVG[spec['chart_type']](spec, largura, altura)
    titulo = spec['title']
    rodape = f"Fonte: {spec.get('source') or '—'} · {spec.get('freshness') or ''} · {spec.get('confidence')}"
    return (
        f'<svg viewBox="0 0 {largura} {altura}" xmlns="http://www.w3.org/2000/svg" role="img" '
        f'aria-label="{titulo}">'
        f'<rect width="{largura}" height="{altura}" fill="{_COR_FUNDO}" />'
        f'<text x="20" y="30" fill="{_COR_DOURADO}" font-size="16" font-weight="bold">{titulo}</text>'
        f'{corpo}'
        f'<text x="20" y="{altura - 8}" fill="{_COR_MUTED}" font-size="10">{rodape}</text>'
        f'</svg>'
    )


def adicionar_grafico_ao_slide(slide, spec, left=Inches(0.9), top=Inches(1.7), width=Inches(11.5), height=Inches(5.0)):
    """Gráfico NATIVO do PowerPoint (editável, sem rasterização) --
    mesma especificação canônica usada em renderizar_svg, outro
    renderizador. Nunca chamado com NOT_ENOUGH_DATA -- quem convoca
    decide o slide de 'aguardando dados' separadamente."""
    validar_chart_spec(spec)
    dados = CategoryChartData()
    dados.categories = spec['x']
    for serie in spec['series']:
        dados.add_series(serie['name'], serie['values'])
    grafico_shape = slide.shapes.add_chart(_TIPO_PPTX[spec['chart_type']], left, top, width, height, dados)
    grafico = grafico_shape.chart
    grafico.has_title = True
    grafico.chart_title.text_frame.text = spec['title']
    return grafico


def gerar_e_registrar_grafico(factory, spec, *, meeting_id=None, agent_id=None, decision_id=None,
                               parent_artifact_id=None):
    """Persiste o gráfico como artefato CHART (SVG) via mi_artefatos (M1).
    Nunca chamado com NOT_ENOUGH_DATA -- o chamador decide separadamente
    o que mostrar quando não há dado suficiente."""
    conteudo = renderizar_svg(spec).encode('utf-8')
    return registrar_artefato(
        factory, artifact_type='CHART', conteudo=conteudo, mime_type=MIME_SVG,
        source_type=spec.get('source'), meeting_id=meeting_id, agent_id=agent_id,
        decision_id=decision_id, parent_artifact_id=parent_artifact_id,
        metadata={'chart_type': spec['chart_type'], 'title': spec['title'], 'units': spec.get('units'),
                  'confidence': spec.get('confidence'), 'freshness': spec.get('freshness')},
    )


def registrar_rotas(app, factory, autorizado):
    """Só recebe um chart_spec JÁ MONTADO com dados reais por quem chama
    -- este módulo nunca busca dado de negócio sozinho (isso pertenceria
    a quem consulta CRM/Relationship 360/etc., não ao Chart Engine)."""
    from flask import jsonify, request

    @app.route('/api/admin/mi/graficos', methods=['POST'])
    def mi_chart_gerar():
        if not autorizado():
            return jsonify(success=False, error='Não autorizado.'), 401
        corpo = request.get_json(silent=True) or {}
        spec = corpo.get('spec')
        if spec == NOT_ENOUGH_DATA:
            return jsonify(success=False, error='Dados insuficientes para gerar este gráfico.',
                            motivo=NOT_ENOUGH_DATA), 422
        if not isinstance(spec, dict):
            return jsonify(success=False, error='Campo "spec" é obrigatório.'), 400
        try:
            validar_chart_spec(spec)
        except ValueError as erro:
            return jsonify(success=False, error=str(erro)), 400
        try:
            resultado = gerar_e_registrar_grafico(
                factory, spec, meeting_id=corpo.get('meeting_id'), agent_id=corpo.get('agent_id'),
                decision_id=corpo.get('decision_id'), parent_artifact_id=corpo.get('parent_artifact_id'),
            )
        except Exception:
            app.logger.exception('Falha ao gerar/persistir gráfico')
            return jsonify(success=False, error='Não foi possível gerar o gráfico.'), 503
        sucesso = resultado.pop('success', False)
        return jsonify(success=sucesso, **resultado), (201 if sucesso else 503)
