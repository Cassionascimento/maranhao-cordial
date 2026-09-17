"""P5.X — Milestone 5: Chart Engine. Gráficos SEMPRE programáticos a
partir de números reais; NOT_ENOUGH_DATA explícito quando não há dado
suficiente, nunca um gráfico fabricado. Testa o renderizador SVG (texto
puro, sem dependência nova), o gráfico nativo do PPTX (reaberto pela
própria biblioteca) e a persistência via mi_artefatos."""
import io
import unittest
from unittest.mock import MagicMock

from flask import Flask
from pptx import Presentation
from pptx.util import Inches

import mi_chart_engine as chart
from tests.test_mi_artefatos import FakeArtefatoConn, _db_vazio


def _valor(v):
    return v.adapted if hasattr(v, 'adapted') else v


def _spec(**over):
    base = {
        'chart_type': 'bar', 'title': 'Oportunidades por tipo', 'x': ['recompra', 'upsell'],
        'series': [{'name': 'contagem', 'values': [3, 2]}], 'units': 'unidades',
        'source': 'mi_intelligence_api.overview', 'freshness': '2026-09-17', 'confidence': 'REAL',
    }
    base.update(over)
    return base


class ValidarChartSpec(unittest.TestCase):
    def test_spec_valida_passa(self):
        chart.validar_chart_spec(_spec())  # não deve levantar

    def test_chart_type_invalido_recusa(self):
        with self.assertRaises(ValueError):
            chart.validar_chart_spec(_spec(chart_type='radar'))

    def test_serie_com_tamanho_diferente_do_eixo_x_recusa(self):
        with self.assertRaises(ValueError):
            chart.validar_chart_spec(_spec(series=[{'name': 'c', 'values': [1, 2, 3]}]))

    def test_confidence_fora_do_vocabulario_recusa(self):
        with self.assertRaises(ValueError):
            chart.validar_chart_spec(_spec(confidence='TALVEZ'))

    def test_sem_titulo_recusa(self):
        with self.assertRaises(ValueError):
            chart.validar_chart_spec(_spec(title='  '))

    def test_sem_x_recusa(self):
        with self.assertRaises(ValueError):
            chart.validar_chart_spec(_spec(x=[]))


class MontarChartSpecOuAguardando(unittest.TestCase):
    def test_dados_suficientes_devolve_spec_real(self):
        resultado = chart.montar_chart_spec_ou_aguardando(
            chart_type='bar', title='t', x=['a', 'b'], series=[{'name': 's', 'values': [1, 2]}],
            units='u', source='crm', freshness='hoje', confidence='REAL',
        )
        self.assertIsInstance(resultado, dict)

    def test_confidence_not_enough_data_devolve_marcador_nunca_spec(self):
        resultado = chart.montar_chart_spec_ou_aguardando(
            chart_type='bar', title='t', x=['a'], series=[{'name': 's', 'values': [1]}],
            units='u', source='territorio', freshness=None, confidence='NOT_ENOUGH_DATA',
        )
        self.assertEqual(resultado, chart.NOT_ENOUGH_DATA)

    def test_eixo_x_vazio_devolve_marcador_nunca_fabrica_ponto(self):
        resultado = chart.montar_chart_spec_ou_aguardando(
            chart_type='bar', title='t', x=[], series=[], units='u', source='forecast',
            freshness=None, confidence='REAL',
        )
        self.assertEqual(resultado, chart.NOT_ENOUGH_DATA)

    def test_menos_pontos_que_o_minimo_exigido_devolve_marcador(self):
        resultado = chart.montar_chart_spec_ou_aguardando(
            chart_type='line', title='forecast', x=['jan'], series=[{'name': 'vendas', 'values': [10]}],
            units='R$', source='forecast_readiness', freshness=None, confidence='DERIVED', minimo_pontos=6,
        )
        self.assertEqual(resultado, chart.NOT_ENOUGH_DATA)


class RenderizarSVG(unittest.TestCase):
    def test_svg_bar_e_um_svg_valido_com_titulo_e_rotulos(self):
        svg = chart.renderizar_svg(_spec())
        self.assertTrue(svg.startswith('<svg'))
        self.assertIn('Oportunidades por tipo', svg)
        self.assertIn('recompra', svg)
        self.assertIn('upsell', svg)
        self.assertTrue(svg.rstrip().endswith('</svg>'))

    def test_svg_line_com_multiplas_series(self):
        spec = _spec(chart_type='line', x=['jan', 'fev', 'mar'],
                     series=[{'name': 'a', 'values': [1, 2, 3]}, {'name': 'b', 'values': [3, 2, 1]}])
        svg = chart.renderizar_svg(spec)
        self.assertIn('<polyline', svg)
        self.assertEqual(svg.count('<polyline'), 2)

    def test_svg_pie_soma_100_por_cento_visualmente_via_fatias(self):
        spec = _spec(chart_type='pie', x=['A', 'B'], series=[{'name': 'c', 'values': [30, 70]}])
        svg = chart.renderizar_svg(spec)
        self.assertIn('<path', svg)

    def test_fonte_e_confianca_aparecem_no_rodape_nunca_escondidas(self):
        svg = chart.renderizar_svg(_spec())
        self.assertIn('mi_intelligence_api.overview', svg)
        self.assertIn('REAL', svg)

    def test_svg_de_not_enough_data_nunca_e_gerado_diretamente(self):
        with self.assertRaises(ValueError):
            chart.renderizar_svg(chart.NOT_ENOUGH_DATA)


class GraficoNativoNoPptx(unittest.TestCase):
    def test_grafico_e_reaberto_pela_propria_biblioteca(self):
        apresentacao = Presentation()
        slide = apresentacao.slides.add_slide(apresentacao.slide_layouts[6])
        chart.adicionar_grafico_ao_slide(slide, _spec(), left=Inches(1), top=Inches(1), width=Inches(8), height=Inches(5))
        buffer = io.BytesIO()
        apresentacao.save(buffer)
        reaberta = Presentation(io.BytesIO(buffer.getvalue()))
        graficos = [f for f in reaberta.slides[0].shapes if f.has_chart]
        self.assertEqual(len(graficos), 1)
        self.assertEqual(graficos[0].chart.chart_title.text_frame.text, 'Oportunidades por tipo')


class GerarERegistrarGrafico(unittest.TestCase):
    def test_persiste_como_artefato_chart_via_mi_artefatos(self):
        db = _db_vazio()
        resultado = chart.gerar_e_registrar_grafico(lambda: FakeArtefatoConn(db), _spec())
        self.assertTrue(resultado['success'])
        artefato = db['artefatos'][resultado['id']]
        self.assertEqual(_valor(artefato['artifact_type']), 'CHART')
        self.assertEqual(_valor(artefato['mime_type']), chart.MIME_SVG)

    def test_conteudo_persistido_e_svg_valido(self):
        db = _db_vazio()
        chart.gerar_e_registrar_grafico(lambda: FakeArtefatoConn(db), _spec())
        blob = next(iter(db['blobs'].values()))
        conteudo = _valor(blob['conteudo']).decode('utf-8')
        self.assertTrue(conteudo.startswith('<svg'))


class RotaHTTP(unittest.TestCase):
    def _app(self, db, autorizado=lambda: True):
        app = Flask(__name__)
        chart.registrar_rotas(app, lambda: FakeArtefatoConn(db), autorizado)
        return app

    def test_exige_autenticacao(self):
        db = _db_vazio()
        resp = self._app(db, autorizado=lambda: False).test_client().post('/api/admin/mi/graficos', json={'spec': _spec()})
        self.assertEqual(resp.status_code, 401)

    def test_not_enough_data_devolve_422_nunca_fabrica_grafico(self):
        db = _db_vazio()
        resp = self._app(db).test_client().post('/api/admin/mi/graficos', json={'spec': 'NOT_ENOUGH_DATA'})
        self.assertEqual(resp.status_code, 422)
        self.assertEqual(resp.get_json()['motivo'], 'NOT_ENOUGH_DATA')
        self.assertEqual(db['artefatos'], {})

    def test_spec_invalida_e_400(self):
        db = _db_vazio()
        resp = self._app(db).test_client().post('/api/admin/mi/graficos', json={'spec': {'chart_type': 'radar'}})
        self.assertEqual(resp.status_code, 400)

    def test_spec_valida_gera_e_persiste_201(self):
        db = _db_vazio()
        resp = self._app(db).test_client().post('/api/admin/mi/graficos', json={'spec': _spec()})
        self.assertEqual(resp.status_code, 201)
        self.assertTrue(resp.get_json()['success'])
        self.assertEqual(len(db['artefatos']), 1)


if __name__ == '__main__':
    unittest.main()
