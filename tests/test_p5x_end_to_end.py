"""P5X integração local: SQL em memória, zero rede, dados explicitamente sintéticos."""
import io
import unittest
from unittest.mock import patch
from uuid import uuid4
from xml.etree import ElementTree
from zipfile import ZipFile
from flask import Flask
from pptx import Presentation
import mi_conselho as conselho
import mi_secretario_executivo as secretario
import mi_presentation_engine as deck
import mi_chart_engine as chart
import mi_artefatos as artifacts
import mi_label_studio as label
import mi_pirret_criativo as pirret
from mi_image_provider import MockImageProvider
from tests.test_mi_artefatos import FakeArtefatoCursor,FakeArtefatoConn,_db_vazio,_valor
from tests.test_mi_pirret_criativo import _cliente_mock
from tests.test_mi_secretario_executivo import _registro

class Cursor(FakeArtefatoCursor):
    def execute(self,sql,params=()):
        if sql.startswith('INSERT INTO mi_conselho_registros'):
            ident,chave,digest,*vals=params
            row=dict(zip(conselho.CAMPOS_REGISTRO,map(_valor,vals)))
            row.update(id=ident,chave=chave,payload_hash=digest,versao=1,registro_anterior_id=None,criado_em=None)
            if chave in self.db['registros']:self._resultado=None
            else:self.db['registros'][chave]=row;self._resultado={'id':ident}
        elif sql.startswith('SELECT id, payload_hash FROM mi_conselho_registros'):
            self._resultado=self.db['registros'].get(params[0])
        elif sql.startswith('SELECT * FROM mi_conselho_registros WHERE id='):
            self._resultado=next((r for r in self.db['registros'].values() if r['id']==params[0]),None)
        elif sql.startswith('SELECT * FROM mi_artefatos_auditoria'):
            self._resultado=[r for r in self.db['auditoria'] if r['artefato_id']==params[0]]
        else:super().execute(sql,params)
class Conn(FakeArtefatoConn):
    def cursor(self,**kw):return Cursor(self.db)

class EndToEnd(unittest.TestCase):
    def setUp(self):
        self.db=_db_vazio();self.db['registros']={};self.factory=lambda:Conn(self.db)
        no_network=patch('socket.socket.connect',side_effect=AssertionError('rede proibida'))
        no_network.start();self.addCleanup(no_network.stop)

    def test_conselho_secretario_pptx_chart_pirret_storage_http_aprovacao(self):
        body=_registro(chave=str(uuid4()),tipo='reuniao',demanda='SYNTHETIC_TEST — exercício de governança',
            conclusao='SYNTHETIC_TEST — AGUARDANDO DADOS',dados_apresentados={},recomendacoes=[])
        persisted,status=conselho.registrar_registro(self.factory,body)
        self.assertEqual(status,201)
        retry,status=conselho.registrar_registro(self.factory,body);self.assertEqual(status,200)
        self.assertEqual(retry['id'],persisted['id'])
        record=conselho.buscar_registro(Cursor(self.db),persisted['id'])
        with patch.object(secretario,'OpenAI',side_effect=RuntimeError('provider desabilitado')):
            ata=secretario.montar_ata_executiva(record)
        spec={'chart_type':'bar','title':'SYNTHETIC_TEST — contagens de teste','x':['a','b'],
              'series':[{'name':'teste','values':[2,1]}],'units':'fixtures',
              'source':'fixture local','freshness':'teste','confidence':'SYNTHETIC_TEST'}
        a=deck.gerar_e_registrar_apresentacao(self.factory,ata,graficos=[spec,chart.NOT_ENOUGH_DATA])
        g=chart.gerar_e_registrar_grafico(self.factory,spec,meeting_id=persisted['id'])
        with patch('mi_brand_context.montar_brand_context_completo',return_value={'referencias_aprovadas':[]}):
            visual=label.gerar_conceito_rotulo(self.factory,'SYNTHETIC_TEST','TEST',quantidade=1,provider=MockImageProvider(),cliente=_cliente_mock())
        vid=visual['artefatos'][0]['id']
        self.assertEqual(self.db['artefatos'][vid]['metadata']['confidence'],'SYNTHETIC_TEST')
        app=Flask(__name__);artifacts.registrar_rotas(app,self.factory,lambda:True);client=app.test_client()
        listed=client.get('/api/admin/mi/artefatos').json['artefatos'];self.assertEqual(len(listed),3)
        downloaded=client.get(f"/api/admin/mi/artefatos/{a['id']}/download")
        self.assertEqual(downloaded.headers['X-Content-Type-Options'],'nosniff')
        pptx=downloaded.data
        with ZipFile(io.BytesIO(pptx)) as z:
            self.assertIsNone(z.testzip())
            for name in z.namelist():
                if name.endswith('.xml'):ElementTree.fromstring(z.read(name))
        presentation=Presentation(io.BytesIO(pptx));self.assertLessEqual(len(presentation.slides),7)
        text=' '.join(s.text for slide in presentation.slides for s in slide.shapes if s.has_text_frame)
        self.assertIn('SYNTHETIC_TEST',text)
        self.assertIn('RECOMENDAÇÕES PARA DECISÃO',text)
        self.assertTrue(any(s.has_chart for slide in presentation.slides for s in slide.shapes))
        for ident,decision in ((vid,'aprovar'),(g['id'],'rejeitar')):
            self.assertEqual(client.post(f'/api/admin/mi/artefatos/{ident}/{decision}',json={'ator':'teste humano'}).status_code,200)
        newer=pirret.refinar_conceito_visual(self.factory,vid,'SYNTHETIC_TEST segunda versão',provider=MockImageProvider())
        chain=client.get(f"/api/admin/mi/artefatos/{newer['id']}/linhagem").json['linhagem']
        self.assertEqual([r['version'] for r in chain],[1,2]);self.assertEqual(chain[1]['status'],'gerado')
        self.assertTrue(client.get(f'/api/admin/mi/artefatos/{vid}/auditoria').json['auditoria'])

    def test_svg_escapa_conteudo_e_rejeita_valores_invalidos(self):
        spec={'chart_type':'bar','title':'<script>alert(1)</script>','x':['<img onerror="x">'],
            'series':[{'name':'teste','values':[1]}],'confidence':'REAL','source':'<svg/onload=x>'}
        svg=chart.renderizar_svg(spec);ElementTree.fromstring(svg)
        self.assertNotIn('<script>',svg);self.assertNotIn('<img',svg)
        for invalid in (None,True,float('nan'),float('inf'),-1):
            spec['series'][0]['values']=[invalid]
            with self.assertRaises(ValueError):chart.validar_chart_spec(spec)

    def test_auditoria_exige_autenticacao(self):
        app=Flask(__name__)
        artifacts.registrar_rotas(app,lambda:self.fail('banco'),lambda:False)
        self.assertEqual(app.test_client().get('/api/admin/mi/artefatos/id/auditoria').status_code,401)

class RotasEGovernanca(unittest.TestCase):
    def test_rotas_p5x_montadas_e_todas_protegidas(self):
        import importlib
        modules=['mi_artefatos','mi_brand_context','mi_chart_engine','mi_presentation_engine',
                 'mi_pirret_criativo','mi_label_studio','mi_social_studio','mi_regulatorio_produto']
        app=Flask(__name__)
        def forbidden():raise AssertionError('sem autenticação não abre banco')
        for name in modules:importlib.import_module(name).registrar_rotas(app,forbidden,lambda:False)
        secretario.registrar_rotas_leitura(app,forbidden,lambda:False)
        client=app.test_client()
        import re
        rules=list(app.url_map.iter_rules())
        self.assertGreater(len(rules),20)
        for rule in rules:
            if not rule.rule.startswith('/api/admin/mi/'):continue
            path=re.sub(r'<[^>]+>','teste',rule.rule)
            for method in rule.methods-{'HEAD','OPTIONS'}:
                self.assertEqual(client.open(path,method=method,json={}).status_code,401,(path,method))

    def test_migrations_sao_aditivas_e_numeracao_unica(self):
        from pathlib import Path
        import re
        paths=list(Path('migrations').glob('*.sql'))
        numbers=[p.name.split('_')[0] for p in paths]
        self.assertEqual(len(numbers),len(set(numbers)))
        for number in ('019','020','021'):
            sql=next(p for p in paths if p.name.startswith(number)).read_text()
            sql=re.sub(r'--[^\n]*','',sql)
            self.assertNotRegex(sql.upper(),r'\b(DELETE|DROP|TRUNCATE)\b')
            self.assertIn('CREATE TABLE IF NOT EXISTS',sql)
