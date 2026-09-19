"""P5.X — Milestone 4: Presentation Engine. Verifica que o .pptx gerado é
um arquivo REAL e válido (reaberto pela própria python-pptx), com 3 a 5
slides por padrão, texto oriundo apenas da ata (nunca fabricado), e que a
persistência via mi_artefatos funciona ponta a ponta com um banco fake em
memória (nunca banco real)."""
import io
import unittest
from unittest.mock import MagicMock
from uuid import uuid4

from flask import Flask
from pptx import Presentation

import mi_presentation_engine as engine
from tests.test_mi_artefatos import FakeArtefatoConn, _db_vazio


def _valor(v):
    return v.adapted if hasattr(v, 'adapted') else v


def _ata(**over):
    base = {
        'meeting_id': 'reg-1', 'titulo': 'Conselho — pricing do Bacuri',
        'narrativa': 'Margem caiu 3pp; Standard recomenda revisar tabela de preço.',
        'participantes': ['standard', 'iris'],
        'decisoes': ['standard: revisar tabela de preço'],
        'proximos_passos': ['revisar tabela de preço'],
        'visuais_sugeridos': [], 'precisa_diretor': False,
    }
    base.update(over)
    return base


def _texto_de_todos_os_slides(apresentacao):
    textos = []
    for slide in apresentacao.slides:
        for forma in slide.shapes:
            if forma.has_text_frame:
                textos.append(forma.text_frame.text)
    return textos


class MontarDeckExecutivo(unittest.TestCase):
    def test_pptx_gerado_e_um_arquivo_real_reabrivel(self):
        conteudo = engine.montar_deck_executivo(_ata())
        # Se não fosse um .pptx real e válido, Presentation() levantaria aqui.
        apresentacao = Presentation(io.BytesIO(conteudo))
        self.assertGreaterEqual(len(apresentacao.slides), 3)

    def test_arquivo_comeca_com_assinatura_zip_ooxml_nunca_html(self):
        conteudo = engine.montar_deck_executivo(_ata())
        self.assertEqual(conteudo[:2], b'PK')  # todo .pptx é um zip OOXML
        self.assertNotIn(b'<html', conteudo[:200].lower())

    def test_padrao_fica_entre_3_e_5_slides(self):
        apresentacao = Presentation(io.BytesIO(engine.montar_deck_executivo(_ata())))
        self.assertGreaterEqual(len(apresentacao.slides), 3)
        self.assertLessEqual(len(apresentacao.slides), 5)

    def test_titulo_e_narrativa_da_ata_aparecem_no_deck(self):
        apresentacao = Presentation(io.BytesIO(engine.montar_deck_executivo(_ata())))
        textos = ' '.join(_texto_de_todos_os_slides(apresentacao))
        self.assertIn('Conselho — pricing do Bacuri', textos)
        self.assertIn('Margem caiu 3pp', textos)

    def test_decisoes_e_proximos_passos_aparecem_no_deck(self):
        apresentacao = Presentation(io.BytesIO(engine.montar_deck_executivo(_ata())))
        textos = ' '.join(_texto_de_todos_os_slides(apresentacao))
        self.assertIn('revisar tabela de preço', textos)

    def test_lista_vazia_mostra_texto_explicito_nunca_omite_o_slide_silenciosamente(self):
        apresentacao = Presentation(io.BytesIO(engine.montar_deck_executivo(_ata(decisoes=[], proximos_passos=[]))))
        textos = ' '.join(_texto_de_todos_os_slides(apresentacao))
        self.assertIn('Nenhuma decisão registrada', textos)
        self.assertIn('Nenhum próximo passo registrado', textos)

    def test_participantes_sao_traduzidos_para_nomes_dos_agentes(self):
        apresentacao = Presentation(io.BytesIO(engine.montar_deck_executivo(_ata(participantes=['standard', 'iris']))))
        textos = ' '.join(_texto_de_todos_os_slides(apresentacao))
        self.assertIn('Standard', textos)
        self.assertIn('Iris', textos)

    def test_mais_de_5_itens_e_cortado_nunca_estoura_o_slide(self):
        muitas_decisoes = [f'decisão {i}' for i in range(10)]
        apresentacao = Presentation(io.BytesIO(engine.montar_deck_executivo(_ata(decisoes=muitas_decisoes))))
        textos = ' '.join(_texto_de_todos_os_slides(apresentacao))
        self.assertIn('decisão 0', textos)
        self.assertNotIn('decisão 9', textos)

    def test_nenhum_html_e_gerado_em_nenhum_lugar_do_modulo(self):
        import inspect
        codigo = inspect.getsource(engine)
        for proibido in ('<html', '<div', '<!DOCTYPE'):
            self.assertNotIn(proibido, codigo)


class DeckComGraficos(unittest.TestCase):
    """P5.X M5 -- wiring do Chart Engine no deck: um slide por gráfico,
    nunca mais de uma ideia dominante por slide, e NOT_ENOUGH_DATA vira
    um slide explícito, nunca um gráfico fabricado."""

    def _spec(self):
        return {
            'chart_type': 'bar', 'title': 'Oportunidades por tipo', 'x': ['recompra', 'upsell'],
            'series': [{'name': 'contagem', 'values': [3, 2]}], 'units': 'unidades',
            'source': 'mi_intelligence_api.overview', 'freshness': '2026-09-17', 'confidence': 'REAL',
        }

    def test_grafico_real_vira_slide_com_grafico_nativo_do_pptx(self):
        apresentacao = Presentation(io.BytesIO(engine.montar_deck_executivo(_ata(), graficos=[self._spec()])))
        self.assertEqual(len(apresentacao.slides), 5)  # capa+narrativa+grafico+decisoes+proximos
        graficos = [f for slide in apresentacao.slides for f in slide.shapes if f.has_chart]
        self.assertEqual(len(graficos), 1)

    def test_not_enough_data_vira_slide_explicito_nunca_grafico_fabricado(self):
        import mi_chart_engine
        apresentacao = Presentation(io.BytesIO(
            engine.montar_deck_executivo(_ata(), graficos=[mi_chart_engine.NOT_ENOUGH_DATA])
        ))
        textos = ' '.join(_texto_de_todos_os_slides(apresentacao))
        self.assertIn('Aguardando dados suficientes', textos)
        graficos = [f for slide in apresentacao.slides for f in slide.shapes if f.has_chart]
        self.assertEqual(len(graficos), 0)

    def test_sem_graficos_deck_permanece_com_4_slides(self):
        apresentacao = Presentation(io.BytesIO(engine.montar_deck_executivo(_ata())))
        self.assertEqual(len(apresentacao.slides), 4)


class GerarERegistrarApresentacao(unittest.TestCase):
    def test_persiste_como_artefato_presentation_via_mi_artefatos(self):
        db = _db_vazio()
        resultado = engine.gerar_e_registrar_apresentacao(lambda: FakeArtefatoConn(db), _ata())
        self.assertTrue(resultado['success'])
        artefato = db['artefatos'][resultado['id']]
        self.assertEqual(_valor(artefato['artifact_type']), 'PRESENTATION')
        self.assertEqual(_valor(artefato['mime_type']), engine.MIME_PPTX)

    def test_conteudo_persistido_e_um_pptx_reabrivel(self):
        db = _db_vazio()
        engine.gerar_e_registrar_apresentacao(lambda: FakeArtefatoConn(db), _ata())
        blob = next(iter(db['blobs'].values()))
        conteudo = _valor(blob['conteudo'])
        Presentation(io.BytesIO(conteudo))  # não deve levantar


class RotaHTTP(unittest.TestCase):
    def _registro(self):
        return {
            'id': 'reg-1', 'chave': 'k1', 'versao': 1, 'criado_em': None,
            'tipo': 'reuniao', 'demanda': 'pricing do Bacuri', 'participantes': ['standard'],
            'contexto': None, 'dados_apresentados': {
                'sintese_estruturada': {'proxima_acao': 'revisar preço'},
                'pareceres_compactos': [{'agente': 'standard', 'requested_visual': None}],
            },
            'posicoes': None, 'conflitos': None, 'conclusao': 'ajustar pricing',
            'recomendacoes': [{'responsavel': 'standard', 'descricao': 'revisar preço'}],
            'vetos': None, 'pendencias': None, 'precisa_diretor': False,
        }

    def _app(self, db, registro, autorizado=lambda: True):
        app = Flask(__name__)
        conn = MagicMock()
        conn.cursor.return_value.__enter__.return_value.fetchone.return_value = registro
        # a rota abre duas "conexões" -- uma de leitura (mock acima) e,
        # dentro de gerar_e_registrar_apresentacao, uma de escrita via a
        # MESMA factory (aqui simplificada: a leitura usa o mock, a
        # escrita usa o fake in-memory real de mi_artefatos).
        chamadas = {'n': 0}

        def factory():
            chamadas['n'] += 1
            return conn if chamadas['n'] == 1 else FakeArtefatoConn(db)

        engine.registrar_rotas(app, factory, autorizado)
        return app

    def test_exige_autenticacao(self):
        db = _db_vazio()
        app = self._app(db, None, autorizado=lambda: False)
        resp = app.test_client().post('/api/admin/mi/conselho/reg-1/apresentacao')
        self.assertEqual(resp.status_code, 401)

    def test_registro_inexistente_e_404(self):
        db = _db_vazio()
        app = self._app(db, None)
        resp = app.test_client().post('/api/admin/mi/conselho/reg-1/apresentacao')
        self.assertEqual(resp.status_code, 404)

    def test_gera_e_persiste_via_http_devolve_id_do_artefato(self):
        db = _db_vazio()
        app = self._app(db, self._registro())
        resp = app.test_client().post('/api/admin/mi/conselho/reg-1/apresentacao')
        self.assertEqual(resp.status_code, 201)
        corpo = resp.get_json()
        self.assertTrue(corpo['success'])
        self.assertIn('id', corpo)
        self.assertEqual(len(db['artefatos']), 1)


if __name__ == '__main__':
    unittest.main()
